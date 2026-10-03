"""Test cho tính năng nhập kịch bản từ CSV/Excel (đã build vòng 4)."""
import io

import pytest

from app.pipeline.script_import import ScriptImportError, parse_script_rows, parse_ts

HEADER = ["Mã block", "Thời lượng", "Loại Visual", "Hình ảnh & Hiệu ứng (Visual/FX)", "Âm thanh & Nhạc nền (Audio/SFX)", "Kịch bản Giọng đọc (VO Content)"]


def _csv_bytes(rows: list[list[str]]) -> bytes:
    lines = [",".join(f'"{c}"' for c in row) for row in rows]
    return ("\n".join(lines)).encode("utf-8")


# ---------------------------------------------------------------------------
# Unit test thuần cho module parser
# ---------------------------------------------------------------------------
@pytest.mark.parametrize(
    "ts,expected",
    [
        ("0:00–0:05", (0, 5)),
        ("0:05-0:20", (5, 20)),
        ("1:30", (90, None)),
        ("không hợp lệ", (None, None)),
        ("", (None, None)),
    ],
)
def test_parse_ts(ts, expected):
    assert parse_ts(ts) == expected


def test_parse_script_rows_happy_path():
    rows = [
        HEADER,
        ["B01", "0:00–0:05", "Video", "Cinematic opening", "Nhạc nền dồn dập", "Đây là hook mở đầu."],
        ["B02", "0:05-0:20", "Image", "Infographic chart", "Tiếng gõ bàn phím", "Nội dung chính đoạn 2."],
    ]
    result = parse_script_rows(rows)
    assert result["stats"]["block_count"] == 2
    assert result["stats"]["duration_label"] == "0:20"
    assert result["beats"][0]["block_id"] == "B01"
    assert result["beats"][0]["timestamp_sec"] == 0
    assert result["beats"][0]["end_sec"] == 5
    assert result["beats"][0]["direction_label"] == "Audio/SFX"
    assert result["beats"][0]["direction"] == "Nhạc nền dồn dập"
    assert result["beats"][1]["visual_type"] == "Image"
    assert "hook mở đầu" in result["full_text"]


def test_parse_script_rows_column_order_flexible():
    """Cột có thể đổi thứ tự, chỉ cần khớp từ khoá tiếng Việt trong header."""
    header = ["Kịch bản Giọng đọc (VO Content)", "Mã block", "Âm thanh & Nhạc nền (Audio/SFX)", "Thời lượng", "Hình ảnh & Hiệu ứng (Visual/FX)", "Loại Visual"]
    rows = [header, ["Lời thoại", "B01", "SFX", "0:00", "Visual", "Video"]]
    result = parse_script_rows(rows)
    assert result["beats"][0]["block_id"] == "B01"
    assert result["beats"][0]["audio"] == "Lời thoại"


def test_parse_script_rows_missing_column_raises():
    rows = [["Mã block", "Thời lượng"], ["B01", "0:00"]]
    with pytest.raises(ScriptImportError):
        parse_script_rows(rows)


def test_parse_script_rows_empty_raises():
    with pytest.raises(ScriptImportError):
        parse_script_rows([])


def test_parse_script_rows_no_data_rows_raises():
    with pytest.raises(ScriptImportError):
        parse_script_rows([HEADER])


def test_parse_script_rows_unparseable_ts_falls_back_sequential():
    rows = [HEADER, ["B01", "abc", "Image", "v", "d", "audio 1"], ["B02", "xyz", "Image", "v", "d", "audio 2"]]
    result = parse_script_rows(rows)
    ts0 = result["beats"][0]["timestamp_sec"]
    ts1 = result["beats"][1]["timestamp_sec"]
    assert ts1 > ts0  # vẫn tăng dần dù không đọc được thời lượng thật
    assert result["stats"]["duration_label"] == "Không xác định"


# ---------------------------------------------------------------------------
# Test API — parse (preview) + confirm
# ---------------------------------------------------------------------------
def test_import_parse_endpoint_success(client, project):
    csv_bytes = _csv_bytes([HEADER, ["B01", "0:00–0:05", "Video", "Cảnh mở", "Nhạc nền", "Xin chào các bạn."]])
    resp = client.post(f"/projects/{project['id']}/script/import/parse", files={"file": ("script.csv", io.BytesIO(csv_bytes), "text/csv")})
    assert resp.status_code == 200
    data = resp.json()
    assert data["stats"]["block_count"] == 1
    assert data["beats"][0]["audio"] == "Xin chào các bạn."

    # parse KHÔNG được đụng vào pack — project vẫn ở step 0
    proj = client.get(f"/projects/{project['id']}").json()
    assert proj["step"] == 0


def test_import_parse_endpoint_invalid_file_returns_400(client, project):
    resp = client.post(f"/projects/{project['id']}/script/import/parse", files={"file": ("script.csv", io.BytesIO(b"khong,phai,dung,cot"), "text/csv")})
    assert resp.status_code == 400
    assert "5 cột" in resp.json()["detail"]


def test_import_confirm_jumps_straight_to_script_studio(client, project):
    """Import phải nhảy thẳng qua Script Studio (step 1, đã duyệt) mà KHÔNG cần gọi AI
    Generation nào — con đường DUY NHẤT để có script (2026-08-17, mục 44)."""
    pid = project["id"]
    csv_bytes = _csv_bytes(
        [
            HEADER,
            ["B01", "0:00–0:05", "Video", "Cảnh mở màn cinematic", "Nhạc nền dồn dập", "Đây là hook mở đầu rất mạnh."],
            ["B02", "0:05-0:30", "Image", "Infographic số liệu", "Tiếng gõ bàn phím", "Đây là nội dung chính của video."],
        ]
    )
    parsed = client.post(f"/projects/{pid}/script/import/parse", files={"file": ("script.xlsx" if False else "script.csv", io.BytesIO(csv_bytes), "text/csv")}).json()

    resp = client.post(f"/projects/{pid}/script/import/confirm", json={"beats": parsed["beats"], "full_text": parsed["full_text"]})
    assert resp.status_code == 200
    pack = resp.json()
    assert pack["script"]["source"] == "import"
    assert len(pack["script"]["body"]) == 2
    assert pack["script"]["body"][0]["block_id"] == "B01"
    assert pack["retention_check"] is not None  # guardrail vẫn chạy dù không có hook

    proj = client.get(f"/projects/{pid}").json()
    assert proj["step"] == 1
    assert proj["max_step_reached"] >= 1
    assert proj["status"] == "generating"


def test_import_confirm_empty_beats_400(client, project):
    resp = client.post(f"/projects/{project['id']}/script/import/confirm", json={"beats": []})
    assert resp.status_code == 400


# ---------------------------------------------------------------------------
# Giọng đọc đa ngôn ngữ (2026-09-04) — cột VO (XX) thay cho cột VO đơn.
# ---------------------------------------------------------------------------
MULTILANG_HEADER = ["Mã", "Thời lượng", "Loại Visual", "Visual/FX", "Audio/SFX", "VO (VI)", "VO (EN)", "VO (DE)"]


def test_parse_script_rows_multilang_columns():
    rows = [
        MULTILANG_HEADER,
        ["B01", "0:00–0:05", "Video", "Cảnh mở", "Nhạc nền", "Xin chào các bạn.", "Hello everyone.", "Hallo zusammen."],
    ]
    result = parse_script_rows(rows, primary_language="vi")
    beat = result["beats"][0]
    assert beat["audio_by_lang"] == {"vi": "Xin chào các bạn.", "en": "Hello everyone.", "de": "Hallo zusammen."}
    assert beat["audio"] == "Xin chào các bạn."  # audio gốc = primary_language


def test_parse_script_rows_multilang_columns_picks_audio_by_primary_language():
    """`audio` (field gốc, dùng cho render/timestamp) phải LUÔN khớp `primary_language`
    của kênh, không phải cột đầu tiên/cột VI mặc định."""
    rows = [
        MULTILANG_HEADER,
        ["B01", "0:00", "Video", "v", "d", "Xin chào.", "Hello.", "Hallo."],
    ]
    result = parse_script_rows(rows, primary_language="de")
    assert result["beats"][0]["audio"] == "Hallo."


def test_parse_script_rows_multilang_columns_partial_translation_allowed():
    """KHÔNG bắt buộc đủ cả 6 ngôn ngữ — thiếu 1 vài cột vẫn parse được, ô trống coi là
    'chưa dịch' (rỗng), không raise lỗi."""
    header = ["Mã", "Thời lượng", "Loại Visual", "Visual/FX", "Audio/SFX", "VO (VI)", "VO (FR)"]
    rows = [header, ["B01", "0:00", "Video", "v", "d", "Xin chào.", ""]]
    result = parse_script_rows(rows, primary_language="vi")
    assert result["beats"][0]["audio_by_lang"] == {"vi": "Xin chào.", "fr": ""}


def test_parse_script_rows_single_vo_column_still_works_backward_compat():
    """File CŨ (1 cột VO đơn, không có mã ngôn ngữ trong ngoặc) vẫn parse y hệt trước —
    tương thích ngược, không bắt buộc đổi sang cột VO (XX)."""
    rows = [HEADER, ["B01", "0:00", "Image", "v", "d", "Lời thoại cũ."]]
    result = parse_script_rows(rows)
    assert result["beats"][0]["audio"] == "Lời thoại cũ."
    assert result["beats"][0]["audio_by_lang"] == {}


def test_import_parse_endpoint_uses_channel_primary_language(client, channel):
    """`/script/import/parse` phải đọc `BrandProfile.primary_language` của kênh để chọn
    đúng cột VO nào trở thành `audio` gốc."""
    profile = client.get(f"/channels/{channel['id']}/brandprofile").json()
    profile["primary_language"] = "en"
    client.put(f"/channels/{channel['id']}/brandprofile", json=profile)

    proj = client.post(f"/channels/{channel['id']}/projects", json={"title": "Multilang test"}).json()
    csv_bytes = _csv_bytes([MULTILANG_HEADER, ["B01", "0:00", "Video", "v", "d", "Xin chào.", "Hello.", "Hallo."]])
    resp = client.post(f"/projects/{proj['id']}/script/import/parse", files={"file": ("s.csv", io.BytesIO(csv_bytes), "text/csv")})
    assert resp.status_code == 200
    assert resp.json()["beats"][0]["audio"] == "Hello."


def test_build_template_workbook_multilang_has_6_vo_columns():
    from openpyxl import load_workbook

    from app.pipeline.script_import import build_template_workbook

    wb = load_workbook(io.BytesIO(build_template_workbook(multilang=True)))
    header = [c.value for c in next(wb.active.iter_rows(min_row=1, max_row=1))]
    assert header == ["Mã", "Thời lượng", "Loại Visual", "Visual/FX", "Audio/SFX", "VO (VI)", "VO (EN)", "VO (DE)", "VO (PT-BR)", "VO (ES)", "VO (FR)"]


def test_download_script_import_template_multilang_query_param(client, project):
    resp = client.get(f"/projects/{project['id']}/script/import/template", params={"multilang": "true"})
    assert resp.status_code == 200
    assert "da-ngon-ngu" in resp.headers["content-disposition"]


def test_visual_generate_from_imported_script_skips_ai(client, project):
    """Với script nguồn 'import', /visual/generate phải seed shot trực tiếp từ nội
    dung block (không gọi AI diễn giải lại) — xem _seed_shot_from_beat trong pipeline.py."""
    pid = project["id"]
    csv_bytes = _csv_bytes([HEADER, ["B01", "0:00–0:05", "Video", "Mô tả visual chính xác đã viết sẵn", "Mô tả audio chính xác đã viết sẵn", "Lời thoại."]])
    parsed = client.post(f"/projects/{pid}/script/import/parse", files={"file": ("s.csv", io.BytesIO(csv_bytes), "text/csv")}).json()
    client.post(f"/projects/{pid}/script/import/confirm", json={"beats": parsed["beats"], "full_text": parsed["full_text"]})

    resp = client.post(f"/projects/{pid}/visual/generate")
    assert resp.status_code == 200
    shot = resp.json()["shots"][0]
    assert shot["visual_fx"] == "Mô tả visual chính xác đã viết sẵn"
    assert shot["audio_sfx"] == "Mô tả audio chính xác đã viết sẵn"
    assert shot["visual_type"] == "video"
    assert shot["block_id"] == "B01"
